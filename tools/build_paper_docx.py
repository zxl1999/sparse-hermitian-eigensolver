from pathlib import Path
import re
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'PAPER.md'
OUTPUT = ROOT / '论文_稀疏特征值求解方案_草稿.docx'

def font(run, name='Microsoft YaHei', size=10.5, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic

def plain(text):
    return re.sub(r'\[([^\]]+)\]\([^)]*\)', r'\1', text).replace('`', '').replace('**', '')

def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = instruction
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instr, end])

def add_inline(document, text, style=None):
    paragraph = document.add_paragraph(style=style)
    token_re = re.compile(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)')
    cursor = 0
    for match in token_re.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(plain(text[cursor:match.start()])); font(run)
        token = match.group(0)
        if token.startswith('**'):
            run = paragraph.add_run(token[2:-2]); font(run, bold=True)
        elif token.startswith('`'):
            run = paragraph.add_run(token[1:-1]); font(run, name='Consolas', size=9)
        else:
            run = paragraph.add_run(token[1:-1]); font(run, italic=True)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(plain(text[cursor:])); font(run)

def add_table(document, lines):
    rows = []
    for line in lines:
        values = [plain(x.strip()) for x in line.strip().strip('|').split('|')]
        if values and all(re.fullmatch(r':?-{3,}:?', x.replace(' ', '')) for x in values): continue
        rows.append(values)
    if not rows: return
    columns = max(map(len, rows))
    table = document.add_table(rows=len(rows), cols=columns); table.style = 'Table Grid'
    for ri, values in enumerate(rows):
        row = table.rows[ri]
        for ci in range(columns):
            cell = row.cells[ci]; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell.text = ''
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            run = p.add_run(values[ci] if ci < len(values) else ''); font(run, size=8.5, bold=ri == 0)
            if ri == 0:
                shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), 'D9EAF7'); cell._tc.get_or_add_tcPr().append(shading)
        if ri == 0:
            header = OxmlElement('w:tblHeader'); header.set(qn('w:val'), 'true'); row._tr.get_or_add_trPr().append(header)
    document.add_paragraph()

def configure(document):
    section = document.sections[0]; section.top_margin = Cm(2.2); section.bottom_margin = Cm(2.0); section.left_margin = Cm(2.4); section.right_margin = Cm(2.0)
    normal = document.styles['Normal']; normal.font.name = 'Microsoft YaHei'; normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei'); normal.font.size = Pt(10.5); normal.paragraph_format.line_spacing = 1.35; normal.paragraph_format.space_after = Pt(5)
    for name, size, color in [('Heading 1', 15, '1F4E79'), ('Heading 2', 13, '2F75B5'), ('Heading 3', 11.5, '5B9BD5')]:
        style = document.styles[name]; style.font.name = 'Microsoft YaHei'; style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei'); style.font.size = Pt(size); style.font.bold = True; style.font.color.rgb = RGBColor.from_string(color); style.paragraph_format.space_before = Pt(12); style.paragraph_format.space_after = Pt(5); style.paragraph_format.keep_with_next = True
    code = document.styles.add_style('Code', WD_STYLE_TYPE.PARAGRAPH); code.font.name = 'Consolas'; code._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Consolas'); code.font.size = Pt(8.5); code.paragraph_format.line_spacing = 1.0
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER; run = footer.add_run('论文草稿 · 第 '); font(run, size=9); add_field(footer, ' PAGE '); run = footer.add_run(' 页'); font(run, size=9)

def build():
    lines = SOURCE.read_text(encoding='utf-8').splitlines(); document = Document(); configure(document)
    title = lines[0].lstrip('# ').strip(); subtitle = next((plain(x.strip('*')) for x in lines[1:10] if x.startswith('**——')), '')
    metadata = [plain(x[2:].strip()) for x in lines[1:15] if x.startswith('> ')]
    p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Cm(3); run = p.add_run(title); font(run, size=20, bold=True); run.font.color.rgb = RGBColor(31,78,121)
    if subtitle: p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run = p.add_run(subtitle); font(run, size=13, italic=True)
    for item in metadata: p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run = p.add_run(item); font(run, size=10)
    document.add_page_break(); document.add_paragraph('目录', style='Heading 1'); toc = document.add_paragraph(); add_field(toc, ' TOC \\o "1-3" \\h \\z \\u '); document.add_paragraph('打开文档后如目录未自动刷新，请在 Word 中右键目录并选择“更新域”。'); document.add_page_break()
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if not s or s == '---' or s.startswith('# '): i += 1; continue
        if s.startswith('> '): add_inline(document, s[2:], style='Quote'); i += 1; continue
        if s.startswith('```'):
            block = []; i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'): block.append(lines[i]); i += 1
            p = document.add_paragraph(style='Code'); p.paragraph_format.left_indent = Cm(0.5); p.paragraph_format.right_indent = Cm(0.5)
            for n, line in enumerate(block):
                run = p.add_run(line); font(run, name='Consolas', size=8.5)
                if n < len(block)-1: run.add_break()
            i += 1; continue
        if s.startswith('$$'):
            equation = [s.strip('$').strip()]; i += 1
            while i < len(lines) and not lines[i].strip().startswith('$$'): equation.append(lines[i].strip()); i += 1
            if i < len(lines): equation.append(lines[i].strip('$').strip()); i += 1
            p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run = p.add_run(' '.join(x for x in equation if x)); font(run, name='Cambria Math', size=10, italic=True); continue
        heading = re.match(r'^(#{2,4})\s+(.*)$', s)
        if heading:
            level = min(len(heading.group(1))-1, 3); p = document.add_paragraph(style=f'Heading {level}'); run = p.add_run(plain(heading.group(2))); font(run, size={1:15,2:13,3:11.5}[level], bold=True); i += 1; continue
        if s.startswith('|'):
            block = []
            while i < len(lines) and lines[i].strip().startswith('|'): block.append(lines[i]); i += 1
            add_table(document, block); continue
        bullet = re.match(r'^[-*]\s+(.*)$', s); numbered = re.match(r'^\d+[.)]\s+(.*)$', s)
        if bullet or numbered: add_inline(document, (bullet or numbered).group(1), style='List Bullet' if bullet else 'List Number'); i += 1; continue
        paragraph = [s]; i += 1
        while i < len(lines):
            n = lines[i].strip()
            if not n or n == '---' or n.startswith(('#','> ','```','$$','|','- ','* ')) or re.match(r'^\d+[.)]\s+', n): break
            paragraph.append(n); i += 1
        add_inline(document, ' '.join(paragraph))
    document.core_properties.title = title; document.core_properties.subject = '稀疏特征值求解方案论文'; document.core_properties.author = '项目组'; document.core_properties.comments = '由 PAPER.md 生成的 Word 草稿；定稿前需补齐复现材料并更新目录。'; document.save(OUTPUT); print(OUTPUT)

if __name__ == '__main__': build()
